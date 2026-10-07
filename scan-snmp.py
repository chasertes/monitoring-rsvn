import os
import json
import datetime
import ipaddress
from urllib.parse import urlparse
import openpyxl
import asyncio
from concurrent.futures import ThreadPoolExecutor

# КОРРЕКТНЫЙ ИМПОРТ ДЛЯ АКТУАЛЬНОЙ БИБЛИОТЕКИ PySNMP 7.x
from pysnmp.hlapi.v3arch.asyncio import (
    SnmpEngine, CommunityData, UdpTransportTarget, ContextData,
    get_cmd, bulk_cmd
)
from pysnmp.smi.rfc1902 import ObjectType, ObjectIdentity

# НАСТРОЙКИ SNMP И АНАЛИТИКИ
COMMUNITY = 'public'
SNMP_PORT = 161
TIMEOUT = 2.0       
RETRIES = 1
MAX_THREADS = 50  # Количество параллельных потоков (камер одновременно)

# --- БЛОК: ФИЛЬТРАЦИЯ СОБСТВЕННЫХ СЕТЕЙ ---
# 0 = Игнорировать адреса из OurIp (скрывать из отчета)
# 1 = Собирать абсолютно всех клиентов без исключения
ExceptOurNetworks = 0  

# Список ваших IP-адресов или подсетей, которые нужно игнорировать
OurIp = [
    "10.35.2.0/24",  # Будет полностью игнорироваться вся подсеть
    "10.0.70.10"
]
# ------------------------------------------------

SCALAR_METRICS = {
    "sys_descr": "1.3.6.1.2.1.1.1.0",
    "sys_uptime": "1.3.6.1.2.1.1.3.0",
    "sys_name": "1.3.6.1.2.1.1.5.0",
    "ip_in_receives": "1.3.6.1.2.1.4.3.0",
    "ip_in_hdr_errors": "1.3.6.1.2.1.4.4.0",
    "ip_in_addr_errors": "1.3.6.1.2.1.4.5.0",
    "ip_out_requests": "1.3.6.1.2.1.4.10.0",
    "icmp_in_msgs": "1.3.6.1.2.1.5.1.0",
    "icmp_out_echo_reps": "1.3.6.1.2.1.5.22.0",
    "mac_address_v6": "1.3.6.1.2.1.55.1.5.1.8.2",
    "mac_address_v4": "1.3.6.1.2.1.2.2.1.6.2",
    "interface_speed": "1.3.6.1.2.1.2.2.1.5.2"
}

TCP_BASE_OID = (1, 3, 6, 1, 2, 1, 6)

def extract_ip_from_rtsp(rtsp_url):
    if not rtsp_url:
        return None
    try:
        rtsp_url = str(rtsp_url).strip()
        if not rtsp_url.startswith("rtsp://"):
            rtsp_url = "rtsp://" + rtsp_url
        parsed = urlparse(rtsp_url)
        hostname = parsed.hostname
        return hostname.strip() if hostname else None
    except Exception:
        return None

def read_cameras_from_excel(file_path):
    if not os.path.exists(file_path):
        print(f"Ошибка: Файл Excel '{file_path}' не найден!")
        return []
        
    cameras_list = []
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        for row in range(2, sheet.max_row + 1):
            rtsp_link = sheet.cell(row=row, column=4).value
            ip = extract_ip_from_rtsp(rtsp_link)
            if not ip: continue
            
            raw_tz = sheet.cell(row=row, column=1).value
            raw_order = sheet.cell(row=row, column=2).value
            
            tz_str = str(int(raw_tz)) if isinstance(raw_tz, (int, float)) else str(raw_tz or "Н/Д")
            order_str = str(int(raw_order)) if isinstance(raw_order, (int, float)) else str(raw_order or "Н/Д")
                
            camera_meta = {
                "tz_number": tz_str.strip(),
                "order_number": order_str.strip(),
                "rtsp_url": str(rtsp_link).strip(),
                "operator": str(sheet.cell(row=row, column=5).value or "Н/Д").strip(),
                "ip": ip
            }
            cameras_list.append(camera_meta)
        wb.close()
    except Exception as e:
        print(f"Ошибка при чтении Excel: {e}")
    return cameras_list

async def fetch_snmp_scalars_async(ip, engine):
    """Асинсихронный GET запрос по стандартам PySNMP 7.x"""
    results = {}
    try:
        community_data = CommunityData(COMMUNITY, mpModel=1)
        transport = await UdpTransportTarget.create((ip, SNMP_PORT), timeout=TIMEOUT, retries=RETRIES)
        
        error_indication, error_status, error_index, var_binds = await get_cmd(
            engine, community_data, transport, ContextData(),
            *[ObjectType(ObjectIdentity(oid)) for oid in SCALAR_METRICS.values()]
        )
        
        if error_indication or error_status or not var_binds:
            return None
            
        for var_bind in var_binds:
            oid_obj = var_bind[0]
            val_obj = var_bind[1]
            oid_str = str(oid_obj)
            
            from pysnmp.proto.rfc1902 import OctetString
            if isinstance(val_obj, OctetString) and any(mac_oid in oid_str for mac_oid in ["1.3.6.1.2.1.55.1.5.1.8.2", "1.3.6.1.2.1.2.2.1.6.2"]):
                val_str = val_obj.asOctets().hex()
            else:
                val_str = str(val_obj)
            
            for name, target_oid in SCALAR_METRICS.items():
                if oid_str.endswith(target_oid.lstrip('.')):
                    results[name] = val_str
                    break
        return results
    except Exception:
        return None
async def analyze_rtsp_connections_async(ip, engine):
    """Асинхронный Walk по таблице TCP через bulk_cmd (SNMPv2c) с фильтрацией подсетей."""
    rtsp_clients = []
    try:
        community_data = CommunityData(COMMUNITY, mpModel=1)
        transport = await UdpTransportTarget.create((ip, SNMP_PORT), timeout=TIMEOUT, retries=RETRIES)
        current_oid = ObjectType(ObjectIdentity("1.3.6.1.2.1.6"))
        is_walking = True
        
        while is_walking:
            error_indication, error_status, error_index, var_binds_table = await bulk_cmd(
                engine, community_data, transport, ContextData(),
                0, 50, current_oid, lexicographicMode=True
            )
            
            if error_indication or error_status or not var_binds_table:
                break
                
            flat_var_binds = []
            for item in var_binds_table:
                if isinstance(item, (list, tuple)): flat_var_binds.extend(item)
                else: flat_var_binds.append(item)
            
            for var_bind in flat_var_binds:
                try:
                    oid_obj = var_bind[0]
                    val_obj = var_bind[1]
                except Exception:
                    continue
                    
                oid_tuples = tuple(oid_obj)
                state_val = str(val_obj).lower()
                
                if oid_tuples[:7] != TCP_BASE_OID:
                    is_walking = False
                    break
                
                current_oid = ObjectType(ObjectIdentity(oid_tuples))
                
                if '5' not in state_val and 'established' not in state_val:
                    continue

                # Функция-помощник для проверки масок и добавления клиента
                def add_client_if_allowed(remote_ip, remote_port):
                    if remote_ip == "0.0.0.0":
                        return
                    
                    # Проверка подсетей и одиночных IP-адресов
                    if ExceptOurNetworks == 0:
                        try:
                            client_obj = ipaddress.ip_address(remote_ip)
                            for network in OurIp:
                                if "/" not in network:
                                    if client_obj == ipaddress.ip_address(network):
                                        return
                                else:
                                    net_obj = ipaddress.ip_network(network, strict=False)
                                    if client_obj in net_obj:
                                        return
                        except Exception:
                            pass
                            
                    client = {"client_ip": remote_ip, "client_port": int(remote_port)}
                    if client not in rtsp_clients:
                        rtsp_clients.append(client)

                # Вариант 1: Классическая таблица tcpConnTable (.13)
                if len(oid_tuples) == 20 and oid_tuples[7] == 13:
                    local_port = oid_tuples[14]
                    if local_port == 554:
                        remote_ip = f"{oid_tuples[15]}.{oid_tuples[16]}.{oid_tuples[17]}.{oid_tuples[18]}"
                        remote_port = oid_tuples[19]
                        add_client_if_allowed(remote_ip, remote_port)

                # Вариант 2: Современная таблица tcpConnectionTable (.19)
                elif len(oid_tuples) == 24 and oid_tuples[7] == 19:
                    local_port = oid_tuples[16]
                    if local_port == 554:
                        remote_ip = f"{oid_tuples[19]}.{oid_tuples[20]}.{oid_tuples[21]}.{oid_tuples[22]}"
                        remote_port = oid_tuples[23]
                        add_client_if_allowed(remote_ip, remote_port)
                            
    except Exception:
        pass
                    
    return {"active_rtsp_sessions_count": len(rtsp_clients), "connected_clients": rtsp_clients}

async def worker_coroutine(camera_meta, output_dir):
    """Асинхронная корутина сбора данных, запускаемая внутри изолированного потока ОС"""
    ip = camera_meta["ip"]
    engine = SnmpEngine()
    try:
        device_data = await fetch_snmp_scalars_async(ip, engine)
        if not device_data:
            print(f"❌ [ТЗ: {camera_meta['tz_number']}] Камера {ip} недоступна по SNMP.")
            return
            
        rtsp_data = await analyze_rtsp_connections_async(ip, engine)
        device_data.update(rtsp_data)
        device_data.update(camera_meta)
        device_data["scan_timestamp"] = datetime.datetime.now().isoformat()

        output_filename = os.path.join(output_dir, f"{ip}.json")
        with open(output_filename, "w", encoding="utf-8") as json_file:
            json.dump(device_data, json_file, indent=4, ensure_ascii=False)
            
        print(f"✅ [ТЗ: {camera_meta['tz_number']}] Успешно собран IP: {ip}")
    finally:
        engine.close_dispatcher()

def process_camera_in_thread(camera_meta, output_dir):
    """Изолирует асинхронный контекст PySNMP 7 внутри отдельного потока ОС Windows"""
    try:
        asyncio.run(worker_coroutine(camera_meta, output_dir))
    except Exception as e:
        print(f"💥 Ошибка потока для камеры {camera_meta['ip']}: {e}")

def main():
    excel_file = "cameras.xlsx"
    output_dir = "scanned_metrics"
    os.makedirs(output_dir, exist_ok=True)
    
    cameras = read_cameras_from_excel(excel_file)
    if not cameras:
        print("Список камер для опроса пуст.")
        return

    print(f"🚀 Запуск многопоточного опроса {len(cameras)} камер (Потоков: {MAX_THREADS})...")
    print(f"⚙️ Режим фильтрации сетей: {'Игнорируем OurIp подсети' if ExceptOurNetworks == 0 else 'Выключена (собираем всех)'}")
    
    # Нативный пул потоков Windows
    with ThreadPoolExecutor(max_workers=MAX_THREADS) as executor:
        for cam in cameras:
            executor.submit(process_camera_in_thread, cam, output_dir)
            
    print("\n🏁 Глобальное сканирование сети завершено.")

if __name__ == "__main__":
    main()
